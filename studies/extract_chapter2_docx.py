"""Extract proposal text and tables for auditable manuscript development."""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    document = Document(args.source)
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 2
            lines.append(f"{'#' * max(1, min(level, 6))} {text}")
        else:
            lines.append(text)
        lines.append("")

    for table_number, table in enumerate(document.tables, start=1):
        lines.append(f"## Extracted table {table_number}")
        lines.append("")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
        lines.append("")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("\n".join(lines), encoding="utf-8")
    words = sum(len(line.split()) for line in lines)
    print(f"Paragraphs: {len(document.paragraphs)}")
    print(f"Tables: {len(document.tables)}")
    print(f"Inline shapes: {len(document.inline_shapes)}")
    print(f"Extracted words: {words}")
    print(f"Output: {args.output}")


if __name__ == "__main__":
    main()
